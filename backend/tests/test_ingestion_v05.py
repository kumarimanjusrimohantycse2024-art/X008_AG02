from __future__ import annotations

import io
from uuid import uuid4

import pytest
from docx import Document as DocxDocument
from fastapi.testclient import TestClient

from app.core.security import hash_password
from app.db.session import SessionLocal
from app.main import app
from app.models import Job, JobStatus, User, UserRole

client = TestClient(app)
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
PDF_MIME = "application/pdf"


def make_docx(*paragraphs: tuple[str, str | None]) -> bytes:
    document = DocxDocument()
    for text, style in paragraphs:
        document.add_paragraph(text, style=style) if style else document.add_paragraph(text)
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def make_pdf() -> bytes:
    def stream(text: str) -> bytes:
        content = f"BT /F1 12 Tf 40 240 Td ({text}) Tj ET\n".encode()
        return f"<< /Length {len(content)} >>\nstream\n".encode() + content + b"endstream"

    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R 6 0 R] /Count 2 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 300 300] /Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
        stream("EXPERIENCE"),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 300 300] /Resources << /Font << /F1 5 0 R >> >> /Contents 7 0 R >>",
        stream("SKILLS"),
    ]
    output = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, content in enumerate(objects, start=1):
        offsets.append(len(output))
        output.extend(f"{index} 0 obj\n".encode())
        output.extend(content)
        output.extend(b"\nendobj\n")
    xref_offset = len(output)
    output.extend(f"xref\n0 {len(objects) + 1}\n".encode())
    output.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        output.extend(f"{offset:010d} 00000 n \n".encode())
    output.extend(f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_offset}\n%%EOF".encode())
    return bytes(output)


@pytest.fixture
def ingestion_records(tmp_path, monkeypatch: pytest.MonkeyPatch) -> dict[str, object]:
    monkeypatch.setenv("UPLOAD_DIR", str(tmp_path / "uploads"))
    suffix = uuid4().hex
    password = "correct-horse-battery"
    session = SessionLocal()
    owner = User(email=f"ingest-owner-{suffix}@example.test", password_hash=hash_password(password), name="Owner", role=UserRole.RECRUITER)
    other = User(email=f"ingest-other-{suffix}@example.test", password_hash=hash_password(password), name="Other", role=UserRole.RECRUITER)
    session.add_all([owner, other])
    session.flush()
    job = Job(title=f"Ingestion Job {suffix}", description="Test job", status=JobStatus.DRAFT, created_by=owner.id)
    other_job = Job(title=f"Other Ingestion Job {suffix}", description="Test job", status=JobStatus.DRAFT, created_by=other.id)
    session.add_all([job, other_job])
    session.commit()
    records = {"owner": owner, "other": other, "job": job, "other_job": other_job, "password": password, "session": session}
    yield records
    session.delete(job)
    session.delete(other_job)
    session.flush()
    session.delete(owner)
    session.delete(other)
    session.commit()
    session.close()


def token_for(user: User, password: str) -> str:
    response = client.post("/api/auth/login", json={"email": user.email, "password": password})
    assert response.status_code == 200
    return response.json()["access_token"]


def headers(token: str) -> dict[str, str]:
    return {"Authorization": f"Bearer {token}"}


def test_docx_upload_parses_chunks_and_preserves_source_metadata(ingestion_records: dict[str, object]) -> None:
    owner = ingestion_records["owner"]
    token = token_for(owner, ingestion_records["password"])
    resume = make_docx(("EXPERIENCE", "Heading 1"), ("Built Python services for five years.", None), ("SKILLS", "Heading 1"), ("Python, PostgreSQL, and Docker.", None))
    cover = make_docx(("PROFILE", "Heading 1"), ("I build reliable backend systems.", None))
    response = client.post(
        f"/api/jobs/{ingestion_records['job'].id}/applications",
        headers=headers(token),
        data={"candidate_name": "Ada Candidate", "candidate_email": "ADA@Example.Test"},
        files={
            "resume": ("../resume.docx", resume, DOCX_MIME),
            "cover_letter": ("cover.docx", cover, DOCX_MIME),
        },
    )
    assert response.status_code == 201, response.text
    body = response.json()
    assert body["candidate_email"] == "ada@example.test"
    assert body["ingestion_status"] == "COMPLETED"
    assert body["status"] == "COMPLETED"
    assert body["document_count"] == 2
    assert body["chunk_count"] >= 2
    assert all(document["file_name"] != "../resume.docx" for document in body["documents"])
    chunks = client.get(f"/api/applications/{body['id']}/chunks", headers=headers(token))
    assert chunks.status_code == 200
    assert all(chunk["document_id"] for chunk in chunks.json())
    assert all(chunk["page_number"] is None for chunk in chunks.json())


def test_duplicate_document_is_rejected_without_second_copy(ingestion_records: dict[str, object]) -> None:
    owner = ingestion_records["owner"]
    token = token_for(owner, ingestion_records["password"])
    resume = make_docx(("EXPERIENCE", "Heading 1"), ("Python backend development.", None))
    response = client.post(
        f"/api/jobs/{ingestion_records['job'].id}/applications",
        headers=headers(token),
        data={"candidate_name": "Duplicate Candidate"},
        files={"resume": ("resume.docx", resume, DOCX_MIME), "cover_letter": ("cover.docx", resume, DOCX_MIME)},
    )
    assert response.status_code == 422
    assert "Duplicate document" in response.json()["detail"]


def test_pdf_upload_preserves_page_numbers(ingestion_records: dict[str, object]) -> None:
    owner = ingestion_records["owner"]
    token = token_for(owner, ingestion_records["password"])
    response = client.post(
        f"/api/jobs/{ingestion_records['job'].id}/applications",
        headers=headers(token),
        data={"candidate_name": "PDF Candidate"},
        files={"resume": ("resume.pdf", make_pdf(), PDF_MIME)},
    )
    assert response.status_code == 201, response.text
    chunks = client.get(f"/api/applications/{response.json()['id']}/chunks", headers=headers(token)).json()
    assert [chunk["page_number"] for chunk in chunks] == [1, 2]


def test_unsupported_and_malformed_uploads_are_rejected(ingestion_records: dict[str, object]) -> None:
    owner = ingestion_records["owner"]
    token = token_for(owner, ingestion_records["password"])
    response = client.post(
        f"/api/jobs/{ingestion_records['job'].id}/applications",
        headers=headers(token),
        data={"candidate_name": "Bad Candidate"},
        files={"resume": ("resume.exe", b"MZ-not-a-document", "application/octet-stream")},
    )
    assert response.status_code == 422
    assert "Only PDF and DOCX" in response.json()["detail"]


def test_recruiter_cannot_upload_to_or_read_another_jobs_application(ingestion_records: dict[str, object]) -> None:
    owner = ingestion_records["owner"]
    other = ingestion_records["other"]
    owner_token = token_for(owner, ingestion_records["password"])
    other_token = token_for(other, ingestion_records["password"])
    resume = make_docx(("SKILLS", "Heading 1"), ("Python.", None))
    denied = client.post(
        f"/api/jobs/{ingestion_records['other_job'].id}/applications",
        headers=headers(owner_token),
        data={"candidate_name": "Denied Candidate"},
        files={"resume": ("resume.docx", resume, DOCX_MIME)},
    )
    assert denied.status_code == 404
    assert client.get(f"/api/jobs/{ingestion_records['other_job'].id}/applications", headers=headers(owner_token)).status_code == 404
    created = client.post(
        f"/api/jobs/{ingestion_records['other_job'].id}/applications",
        headers=headers(other_token),
        data={"candidate_name": "Other Candidate"},
        files={"resume": ("resume.docx", resume, DOCX_MIME)},
    )
    assert created.status_code == 201
    application_id = created.json()["id"]
    assert client.get(f"/api/applications/{application_id}", headers=headers(owner_token)).status_code == 404
