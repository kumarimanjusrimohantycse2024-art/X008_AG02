from __future__ import annotations

import io
import re
import zipfile
from dataclasses import dataclass

from docx import Document as DocxDocument
from pypdf import PdfReader

from app.models import DocumentType

PDF_MIME = "application/pdf"
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
SUPPORTED_MIMES = {PDF_MIME, DOCX_MIME}
SUPPORTED_EXTENSIONS = {".pdf": DocumentType.RESUME, ".docx": DocumentType.RESUME}
KNOWN_SECTIONS = {
    "SUMMARY", "PROFILE", "EXPERIENCE", "WORK EXPERIENCE", "EMPLOYMENT", "EDUCATION",
    "SKILLS", "PROJECTS", "CERTIFICATIONS", "ACHIEVEMENTS", "TECHNICAL SKILLS",
}


class DocumentValidationError(ValueError):
    pass


@dataclass(frozen=True)
class ParsedSegment:
    text: str
    page_number: int | None
    section: str | None


def normalize_text(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    paragraphs = [re.sub(r"[ \t]+", " ", part).strip() for part in re.split(r"\n{2,}", text)]
    lines = []
    for paragraph in paragraphs:
        if paragraph:
            lines.append(paragraph)
    return "\n\n".join(lines)


def detect_section(text: str) -> str | None:
    candidate = re.sub(r"[^A-Za-z ]", "", text).strip().upper()
    return candidate if candidate in KNOWN_SECTIONS else None


def validate_upload(filename: str, content_type: str | None, data: bytes) -> DocumentType:
    suffix = "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if suffix not in SUPPORTED_EXTENSIONS:
        raise DocumentValidationError("Only PDF and DOCX files are supported.")
    expected_mime = PDF_MIME if suffix == ".pdf" else DOCX_MIME
    if content_type != expected_mime:
        raise DocumentValidationError("The uploaded file type does not match its content type.")
    if suffix == ".pdf":
        if not data.startswith(b"%PDF-"):
            raise DocumentValidationError("The uploaded PDF is malformed.")
        try:
            PdfReader(io.BytesIO(data))
        except Exception as exc:
            raise DocumentValidationError("The uploaded PDF is malformed.") from exc
    else:
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                names = set(archive.namelist())
                if "[Content_Types].xml" not in names or "word/document.xml" not in names:
                    raise DocumentValidationError("The uploaded DOCX is malformed.")
            DocxDocument(io.BytesIO(data))
        except DocumentValidationError:
            raise
        except Exception as exc:
            raise DocumentValidationError("The uploaded DOCX is malformed.") from exc
    return DocumentType.RESUME


def parse_pdf(data: bytes) -> tuple[str, list[ParsedSegment]]:
    reader = PdfReader(io.BytesIO(data))
    segments: list[ParsedSegment] = []
    for index, page in enumerate(reader.pages, start=1):
        text = normalize_text(page.extract_text() or "")
        if text:
            segments.append(ParsedSegment(text=text, page_number=index, section=detect_section(text.split("\n", 1)[0])))
    return "\n\n".join(segment.text for segment in segments), segments


def parse_docx(data: bytes) -> tuple[str, list[ParsedSegment]]:
    document = DocxDocument(io.BytesIO(data))
    segments: list[ParsedSegment] = []
    current_section: str | None = None
    for paragraph in document.paragraphs:
        text = normalize_text(paragraph.text)
        if not text:
            continue
        detected = detect_section(text)
        if detected or paragraph.style.name.lower().startswith("heading"):
            current_section = detected or text.strip()
        segments.append(ParsedSegment(text=text, page_number=None, section=current_section))
    return "\n\n".join(segment.text for segment in segments), segments


def parse_document(document_type: DocumentType, data: bytes) -> tuple[str, list[ParsedSegment]]:
    if data.startswith(b"%PDF-"):
        return parse_pdf(data)
    if data.startswith(b"PK"):
        return parse_docx(data)
    raise DocumentValidationError("Unsupported document content.")
